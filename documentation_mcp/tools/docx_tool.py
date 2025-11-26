from docx import Document
from docx.shared import Inches

class DocxTool:
    def __init__(self):
        self.document = Document()

    def add_title(self, title):
        self.document.add_heading(title, level=1)

    def add_paragraph(self, text):
        self.document.add_paragraph(text)

    def add_image(self, image_path):
        self.document.add_picture(image_path, width=Inches(5))

    def add_table(self, data):
        table = self.document.add_table(rows=len(data), cols=len(data[0]))
        for i, row in enumerate(data):
            for j, cell in enumerate(row):
                table.cell(i, j).text = cell

    def save(self, file_name):
        self.document.save(file_name)

    def add_citation(self, citation):
        self.document.add_paragraph(f"Citation: {citation}")
