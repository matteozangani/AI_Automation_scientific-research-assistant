# PowerPoint Generation Tool

## Features

### Themes
- Supports a variety of themes to enhance the visual appeal of presentations.

### Layouts
- Provides multiple layout options for slides to accommodate different content formats.

### Speaker Notes
- Enables adding speaker notes for each slide to assist during presentations.

### Table Support
- Allows for the inclusion of tables to present data in a structured format.

### Integration with python-pptx
- Utilizes the `python-pptx` library for creating and manipulating PowerPoint files programmatically.

## Example Usage

```python
from pptx import Presentation

# Create a presentation object
presentation = Presentation()

# Add a slide with a title and content layout
slide = presentation.slides.add_slide(presentation.slide_layouts[1])

# Add title and content to the slide
slide.shapes.title.text = "Slide Title"
slide.shapes.placeholders[1].text = "This is the content of the slide."

# Save the presentation
presentation.save('presentation.pptx')
```

## Installation

To use the tool, install `python-pptx` via pip:

```bash
pip install python-pptx
```

## Licenses

This tool is licensed under the MIT License.